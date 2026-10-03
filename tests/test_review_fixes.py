"""Regression checks for exported provenance, workflow serialization and resource ownership."""

import asyncio
import io
import json
from pathlib import Path
import subprocess
import sys
import threading
import zipfile

import httpx
from openpyxl import load_workbook
import pytest

from fact_manager import evidence
from fact_manager.catalog import Catalog
from fact_manager.cli import main
from fact_manager.core import Store, WorkflowBusy, validate_fields
from fact_manager.server import Jobs, make_app
from fact_manager.workflow import (
    HEADERS,
    Ragflow,
    export_review,
    extract,
    import_review,
)
from test_reasoning import approve, derive, setup


@pytest.mark.parametrize("date", ["20261003", "2026-W40-6", "2026-13-01"])
def test_fact_dates_reject_formats_that_break_date_filters(date):
    with pytest.raises(ValueError):
        validate_fields(
            {"entity": "A", "attribute": "B", "value": "C", "valid_from": date}
        )


def test_excel_includes_readonly_derivation_and_check_provenance(tmp_path):
    store = Store(tmp_path / "facts", ["local"])
    premise = approve(store, setup(store))
    derived = derive(store, premise)
    store.save_check(
        derived["id"],
        "needs_review",
        "=This is advisory",
        "reviewer",
        expected_revision=derived["revision"],
        suggestions={"value": "0.34"},
        premise_ids=[premise["id"]],
        premise_revisions={premise["id"]: premise["revision"]},
    )
    path = tmp_path / "review.xlsx"
    assert export_review(store, path) == 1
    book = load_workbook(path, data_only=False)
    sheet = book["核验与推导"]
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert row["fact_id"] == derived["id"] and row["来源类型"] == "推导候选"
    assert "前提原文" in row["引用说明"]
    assert row["检查摘要"] == "=This is advisory"
    assert row["推导过程"] == derived["origin"]["reasoning"]
    assert json.loads(row["推导前提"])[0]["id"] == premise["id"]
    assert all(cell.data_type != "f" for cell in sheet[2])
    # Advisory edits are ignored by import; the facts sheet remains the review contract.
    sheet.cell(2, 11).value = "modified opinion"
    book["facts"].cell(2, HEADERS.index("decision") + 1).value = "approve"
    book.save(path)
    assert import_review(store, path, "administrator") == 1
    assert store.fact(derived["id"])["check"]["summary"] == "=This is advisory"
    assert store.fact(derived["id"])["value"] == "1/3"


def test_published_facts_do_not_include_unapproved_check_suggestions(tmp_path):
    store = Store(tmp_path, ["local"])
    fact = setup(store)
    store.save_check(
        fact["id"],
        "needs_review",
        "Unapproved suggestion",
        "reviewer",
        expected_revision=fact["revision"],
        suggestions={"value": "1/2"},
    )
    approve(store, fact)
    store.publish("administrator")
    assert store.fact(fact["id"])["check"]["suggestions"]["value"] == "1/2"
    published = store.verified()[0]
    assert published["value"] == "2/3" and published["origin"]["kind"] == "direct"
    assert "check" not in published


def test_workflow_lock_cross_process_and_release_on_error(tmp_path):
    store = Store(tmp_path, ["local"])
    other = Store(tmp_path, ["local"])
    with store.workflow_lock():
        with pytest.raises(WorkflowBusy):
            with other.workflow_lock():
                pytest.fail("same-library overlap")
        result = subprocess.run(
            [
                sys.executable,
                "-c",
                """import sys
from fact_manager.core import Store, WorkflowBusy
try:
    with Store(sys.argv[1], ['local']).workflow_lock(): pass
except WorkflowBusy:
    sys.exit(12)
""",
                str(tmp_path),
            ],
            capture_output=True,
            timeout=10,
        )
        assert result.returncode == 12, result.stderr.decode()
        # Another library remains independently available.
        with Store(tmp_path / "other", ["local"]).workflow_lock():
            pass
    with pytest.raises(RuntimeError):
        with store.workflow_lock():
            raise RuntimeError("interrupted")
    with other.workflow_lock():
        pass


def test_all_library_sync_skips_busy_library_and_manual_sync_reports_busy(
    tmp_path, monkeypatch, capsys
):
    config = {"state_dir": str(tmp_path / "state"), "ragflow_url": "http://ragflow"}
    path = tmp_path / "config.json"
    path.write_text(json.dumps(config))
    catalog = Catalog(config)
    lid = catalog.create("Test")["id"]
    monkeypatch.setattr(
        Ragflow, "sync", lambda *a, **kw: pytest.fail("busy library must not sync")
    )
    with catalog.store(lid).workflow_lock():
        monkeypatch.setattr(
            sys, "argv", ["fact-manager", "--config", str(path), "sync"]
        )
        assert main() is None
        assert json.loads(capsys.readouterr().out)["libraries"] == [
            {"id": lid, "skipped": "busy"}
        ]
        monkeypatch.setattr(
            sys,
            "argv",
            ["fact-manager", "--config", str(path), "--library", lid, "sync"],
        )
        assert main() == 1
        assert "retry later" in capsys.readouterr().err
        jobs = Jobs(catalog)
        try:
            job = catalog.create_job(lid, "check")
            jobs.run(job, lid, "check")
            result = catalog.jobs(lid)[0]
            assert result["status"] == "failed" and "retry later" in result["error"]
        finally:
            jobs.executor.shutdown()


def test_native_cache_reuses_structure_without_sharing_mutable_results(
    tmp_path, monkeypatch
):
    store = Store(tmp_path, ["dataset"])
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>original</w:t></w:r></w:p></w:body></w:document>',
        )
    sid = store.register(
        "dataset",
        "doc",
        "test.docx",
        content.getvalue(),
        [{"id": "a", "content": "original"}, {"id": "b", "content": "original"}],
    )
    original = evidence._native_uncached
    calls = []

    def parse(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(evidence, "_native_uncached", parse)
    first = evidence.inspect(store, sid, "a")
    first["native"]["text"] = "mutated"
    first["native"]["warnings"].append("mutated")
    second = evidence.inspect(store, sid, "b")
    assert len(calls) == 1 and second["native"]["text"] == "original"
    assert "mutated" not in second["native"]["warnings"]
    Path(store.source(sid)["original_path"]).write_bytes(b"tampered")
    with pytest.raises(ValueError, match="SHA256"):
        evidence.inspect(store, sid, "a")


def test_native_cache_keeps_pdf_pages_separate_and_local_titles_are_text(tmp_path):
    from pypdf import PdfWriter

    store = Store(tmp_path, ["dataset", "local"])
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.add_blank_page(width=200, height=200)
    content = io.BytesIO()
    pdf.write(content)
    sid = store.register(
        "dataset",
        "pdf",
        "test.pdf",
        content.getvalue(),
        [
            {"id": "a", "content": "OCR A", "positions": [[1]]},
            {"id": "b", "content": "OCR B", "positions": [[2]]},
        ],
    )
    assert evidence.inspect(store, sid, "a")["native"]["pages"][0]["page"] == 1
    assert evidence.inspect(store, sid, "b")["native"]["pages"][0]["page"] == 2
    sid = store.register(
        "local",
        "text",
        "title.pdf",
        "这是录入的文字".encode(),
        [{"id": "c", "content": "这是录入的文字"}],
    )
    inspected = evidence.inspect(store, sid, "c")
    assert (
        inspected["native"]["format"] == "text"
        and inspected["native"]["text"] == "这是录入的文字"
    )
    assert not inspected["native"]["warnings"]


def test_owned_http_clients_close_on_errors_and_borrowed_clients_remain_open(
    tmp_path, monkeypatch
):
    original = httpx.Client
    created = []

    def create_client(**kwargs):
        client = original(
            transport=httpx.MockTransport(lambda request: httpx.Response(503))
        )
        created.append(client)
        return client

    monkeypatch.setattr("fact_manager.workflow.httpx.Client", create_client)
    monkeypatch.setenv("RAGFLOW_API_KEY", "synthetic-key")
    with pytest.raises(httpx.HTTPStatusError):
        with Ragflow({"ragflow_url": "http://ragflow"}) as rag:
            rag.data("GET", "/datasets")
    assert created[-1].is_closed
    store = Store(tmp_path, ["local"])
    setup(store)
    monkeypatch.setenv("LLM_API_KEY", "synthetic-key")
    with pytest.raises(httpx.HTTPStatusError):
        extract(store, {"llm_model": "mock", "llm_url": "http://model"})
    assert created[-1].is_closed
    with original(
        transport=httpx.MockTransport(lambda request: httpx.Response(503))
    ) as borrowed:
        with Ragflow({"ragflow_url": "http://ragflow"}, borrowed):
            pass
        assert not borrowed.is_closed
        with pytest.raises(httpx.HTTPStatusError):
            extract(store, {"llm_model": "mock", "llm_url": "http://model"}, borrowed)
        assert not borrowed.is_closed


def test_slow_dataset_request_does_not_block_another_api_request(tmp_path, monkeypatch):
    config = {
        "state_dir": str(tmp_path),
        "llm_model": "mock",
        "ragflow_url": "http://ragflow",
        "public_host": "facts.example.test",
    }
    catalog = Catalog(config)
    monkeypatch.setenv("FACT_ADMIN_TOKEN", "a" * 16)
    started = threading.Event()
    release = threading.Event()

    def datasets(*args, **kwargs):
        started.set()
        if not release.wait(10):
            raise RuntimeError("test timed out")
        return []

    monkeypatch.setattr(Ragflow, "data", datasets)

    async def requests():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=make_app(catalog, config)),
            base_url="https://facts.example.test",
        ) as client:
            assert (
                await client.post("/api/login", json={"token": "a" * 16})
            ).status_code == 200
            slow = asyncio.create_task(client.get("/api/ragflow/datasets"))
            try:
                assert await asyncio.to_thread(started.wait, 3)
                response = await asyncio.wait_for(client.get("/api/session"), timeout=2)
                assert response.status_code == 200
            finally:
                release.set()
                assert (await slow).status_code == 200

    asyncio.run(requests())


def test_web_excel_import_preserves_review_contract_after_async_upload(catalog, web):
    library_id = catalog.create("Workbook round trip")["id"]
    store = catalog.store(library_id)
    fact = setup(store)
    base = "/api/libraries/" + library_id
    response = web.get(base + "/review.xlsx")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    book["facts"].cell(2, HEADERS.index("decision") + 1).value = "approve"
    content = io.BytesIO()
    book.save(content)
    response = web.post(
        base + "/review-import",
        files={
            "file": (
                "review.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"reviewer": "administrator"},
    )
    assert response.status_code == 200, response.text
    assert response.json() == {"reviewed": 1}
    assert store.fact(fact["id"])["status"] == "approved"
    assert store.verified() == []
