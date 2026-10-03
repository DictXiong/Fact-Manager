import json

import httpx
from openpyxl import load_workbook
import pytest

from fact_manager.core import Store
from fact_manager.workflow import (
    HEADERS,
    Ragflow,
    export_review,
    extract,
    import_review,
)


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path, ["company"])


def source(store, original=b"original-v1", content="研发人员为12人。", document="doc"):
    return store.register(
        "company",
        document,
        "公司介绍.docx",
        original,
        [{"id": "chunk", "content": content, "positions": [[1, 0, 100, 0, 100]]}],
    )


def candidate(value="12", **kwargs):
    return {
        "entity": "公司",
        "attribute": "研发人数",
        "value": value,
        "unit": "人",
        "conditions": "",
        "valid_from": "",
        "valid_until": "",
        "quote": "研发人员为12人。",
        **kwargs,
    }


def approve_all(store, tmp_path, external="yes"):
    path = tmp_path / "review.xlsx"
    export_review(store, path)
    book = load_workbook(path)
    for row in book["facts"].iter_rows(min_row=2):
        row[HEADERS.index("decision")].value = "approve"
        row[HEADERS.index("external_use")].value = external
    book.save(path)
    import_review(store, path, "审核人")
    return path


def test_review_and_publication_are_separate(store, tmp_path):
    sid = source(store)
    assert store.add_candidates(sid, "chunk", [candidate()]) == 1
    assert store.verified() == []
    path = approve_all(store, tmp_path)
    assert store.verified() == []
    assert store.publish("发布人") == 1
    facts = store.verified()
    assert facts[0]["value"] == "12"
    assert facts[0]["quote"] == "研发人员为12人。"
    assert facts[0]["file_hash"]
    with pytest.raises(ValueError, match="already imported"):
        import_review(store, path, "审核人")


def test_external_permission_and_validity_are_enforced(store, tmp_path):
    sid = source(store)
    store.add_candidates(sid, "chunk", [candidate(valid_until="2025-12-31")])
    approve_all(store, tmp_path, external="no")
    store.publish("审核人")
    assert store.verified(on_date="2025-06-01") == []
    assert len(store.verified(for_external=False, on_date="2025-06-01")) == 1
    assert store.verified(for_external=False, on_date="2026-01-01") == []


def test_new_source_version_hides_old_facts_and_stale_review(store, tmp_path):
    sid = source(store)
    store.add_candidates(sid, "chunk", [candidate()])
    path = approve_all(store, tmp_path)
    store.publish("审核人")
    source(store, b"original-v2")
    assert store.verified() == []
    with pytest.raises(ValueError, match="superseded"):
        store.source(sid)
    with pytest.raises(ValueError):
        import_review(store, path, "审核人")


def test_dataset_scope_fails_closed(store, tmp_path):
    sid = source(store)
    store.add_candidates(sid, "chunk", [candidate()])
    approve_all(store, tmp_path)
    store.publish("审核人")
    restricted = Store(store.root, [])
    assert restricted.verified() == []
    assert restricted.sources() == []
    with pytest.raises(ValueError, match="allowlist"):
        restricted.source(sid)


def test_conflicts_block_entire_publication_and_can_be_retired(store, tmp_path):
    sid = source(store)
    store.add_candidates(sid, "chunk", [candidate(), candidate("13")])
    approve_all(store, tmp_path)
    with pytest.raises(ValueError, match="Conflicting"):
        store.publish("审核人")
    assert store.verified() == []
    with store.connect() as db:
        obsolete = db.execute("SELECT id FROM facts WHERE value='13'").fetchone()["id"]
    store.retire(obsolete, "审核人", "手动修正错误提取")
    assert store.publish("审核人") == 1


def test_bad_model_evidence_rolls_back_whole_chunk(store):
    sid = source(store)
    with pytest.raises(ValueError, match="literal substring"):
        store.add_candidates(
            sid, "chunk", [candidate(), candidate("99", quote="不存在的证据")]
        )
    assert store.review_rows() == []


def test_review_cannot_change_evidence(store, tmp_path):
    sid = source(store)
    store.add_candidates(sid, "chunk", [candidate()])
    path = tmp_path / "tampered.xlsx"
    export_review(store, path)
    book = load_workbook(path)
    book["facts"].cell(2, HEADERS.index("decision") + 1).value = "approve"
    book["facts"].cell(2, HEADERS.index("quote") + 1).value = "伪造证据"
    book.save(path)
    with pytest.raises(ValueError, match="Evidence fields"):
        import_review(store, path, "审核人")
    assert len(store.review_rows()) == 1


def test_excel_formulas_are_exported_as_text_and_rejected_on_import(store, tmp_path):
    sid = source(store)
    store.add_candidates(sid, "chunk", [candidate(value="=1+1")])
    path = tmp_path / "formula.xlsx"
    export_review(store, path)
    book = load_workbook(path)
    cell = book["facts"].cell(2, HEADERS.index("value") + 1)
    assert cell.data_type == "s"
    cell.value = "=2+2"
    book.save(path)
    with pytest.raises(ValueError, match="formulas"):
        import_review(store, path, "审核人")


def test_chunk_pagination_reads_entire_document(monkeypatch):
    monkeypatch.setenv("RAGFLOW_API_KEY", "test")
    pages = []

    def reply(request):
        page = int(request.url.params["page"])
        pages.append(page)
        count = 100 if page == 1 else 7
        return httpx.Response(
            200,
            json={
                "code": 0,
                "data": {"chunks": [{"id": f"{page}-{n}"} for n in range(count)]},
            },
        )

    api = Ragflow(
        {"ragflow_url": "http://ragflow"},
        httpx.Client(transport=httpx.MockTransport(reply)),
    )
    assert len(api.chunks("company", "doc")) == 107
    assert pages == [1, 2]


def test_retrieval_rejects_out_of_scope_and_changed_chunks(store, monkeypatch):
    sid = source(store)
    monkeypatch.setenv("RAGFLOW_API_KEY", "test")

    def reply(request):
        body = json.loads(request.content)
        assert body["dataset_ids"] == ["company"]
        assert body["document_ids"] == ["doc"]
        return httpx.Response(
            200,
            json={
                "code": 0,
                "data": {
                    "chunks": [
                        {
                            "id": "chunk",
                            "dataset_id": "company",
                            "document_id": "doc",
                            "content": "研发人员为12人。",
                        },
                        {
                            "id": "chunk",
                            "dataset_id": "company",
                            "document_id": "doc",
                            "content": "重新解析的内容",
                        },
                        {
                            "id": "chunk",
                            "dataset_id": "other",
                            "document_id": "doc",
                            "content": "研发人员为12人。",
                        },
                    ]
                },
            },
        )

    api = Ragflow(
        {"ragflow_url": "http://ragflow"},
        httpx.Client(transport=httpx.MockTransport(reply)),
    )
    result = api.search(store, "研发")
    assert len(result) == 1
    assert result[0]["source_id"] == sid


def test_extraction_is_incremental_and_does_not_publish(store, monkeypatch):
    source(store)
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test")
    calls = []

    def reply(request):
        calls.append(request)
        body = json.loads(request.content)
        assert body["model"] == "deepseek-flash"
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps({"facts": [candidate()]})},
                    }
                ]
            },
        )

    config = {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"}
    client = httpx.Client(transport=httpx.MockTransport(reply))
    assert extract(store, config, client) == 1
    assert extract(store, config, client) == 0
    assert len(calls) == 1
    assert store.verified() == []


def test_extraction_passes_cross_chunk_context_but_rejects_context_only_quotes(
    store, monkeypatch
):
    store.register(
        "company",
        "project",
        "项目介绍.txt",
        b"project",
        [
            {"id": "first", "content": "项目名称：项目甲。张三是项目负责人。"},
            {"id": "second", "content": "工作单位：实验室乙。"},
        ],
    )
    monkeypatch.setenv("LLM_API_KEY", "test")
    seen = []

    def reply(request):
        data = json.loads(request.content)
        supplied = json.loads(data["messages"][1]["content"])
        seen.append(supplied)
        assert "项目名称：项目甲" in supplied["document_context"]["paragraphs"]
        if supplied["target_chunk_index"] == 0:
            assert supplied["next_chunk"] == "工作单位：实验室乙。"
            facts = []
        else:
            assert "张三是项目负责人" in supplied["previous_chunk"]
            # This quote is present in document context, but not in the target chunk.
            facts = [
                candidate(
                    entity="项目甲",
                    attribute="项目负责人",
                    value="张三",
                    unit="",
                    quote="张三是项目负责人。",
                )
            ]
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps({"facts": facts})},
                    }
                ]
            },
        )

    with pytest.raises(ValueError, match="unsupported evidence quote"):
        extract(
            store,
            {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"},
            httpx.Client(transport=httpx.MockTransport(reply)),
        )
    assert len(seen) == 3
    assert store.review_rows() == []


def test_word_context_preserves_person_role_paragraph_order_and_is_bounded(store):
    import io
    import zipfile
    from fact_manager.workflow import document_context

    content = "项目负责人。李四项目商务主管。"
    document = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Li Dan</w:t></w:r></w:p><w:p><w:r><w:t>Principal Investigator</w:t></w:r></w:p><w:p><w:r><w:t>Xiong Dian</w:t></w:r></w:p></w:body></w:document>'
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document)
    sid = source(store, original=buffer.getvalue(), content=content)
    context = document_context(store.source(sid), store.chunks(sid))
    assert context["paragraphs"] == "Li Dan\nPrincipal Investigator\nXiong Dian"
    assert store.chunks(sid)[0]["content"] == content
    context = document_context(
        {"name": "large.txt", "original_path": "/nonexistent"},
        [{"content": "head" + "x" * 100000 + "tail"}],
    )
    assert context["paragraphs"].startswith("head") and context["paragraphs"].endswith(
        "tail"
    )
    assert len(context["paragraphs"]) < 48100


def test_quote_repair_and_prior_facts_survive_incremental_resume(store, monkeypatch):
    sid = store.register(
        "company",
        "document",
        "实验.txt",
        b"original",
        [
            {"id": "first", "content": "研发人员为12人。"},
            {"id": "second", "content": "实验温度为20度。"},
        ],
    )
    monkeypatch.setenv("LLM_API_KEY", "test")
    calls = []
    config = {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"}

    def reply(request):
        data = json.loads(request.content)
        context = json.loads(data["messages"][1]["content"])
        calls.append(data)
        assert data["temperature"] == 0
        if context["target_chunk_index"] == 0:
            assert context["previous_facts"] == []
            facts = [
                candidate(
                    quote=(
                        "上下文中的其他句子"
                        if len(data["messages"]) == 2
                        else "研发人员为12人。"
                    )
                )
            ]
        else:
            assert len(context["previous_facts"]) == 1
            assert context["previous_facts"][0]["value"] == "12"
            facts = []
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps({"facts": facts})},
                    }
                ]
            },
        )

    client = httpx.Client(transport=httpx.MockTransport(reply))
    assert extract(store, config, client) == 1
    assert len(calls) == 3
    assert len(store.review_rows()) == 1
    version = "fact-manager-2|document-context-2"
    assert len(store.extraction_candidates(sid, "deepseek-flash", version)) == 1
    # An interrupted run resumes with prior facts and does not repeat paid completed chunks.
    with store.connect() as db:
        db.execute("DELETE FROM extractions WHERE chunk_id='second'")
    assert extract(store, config, client) == 0
    assert len(calls) == 4
    # A new prompt version must not use old facts to suppress corrected extraction.
    assert store.extraction_candidates(sid, "deepseek-flash", "new-version") == []


def test_identical_extraction_fields_reuse_fact_across_chunks_but_keep_different_conditions(
    store,
):
    sid = store.register(
        "company",
        "doc",
        "实验.txt",
        b"original",
        [
            {"id": "first", "content": "研发人员为12人。"},
            {"id": "second", "content": "研发人数合计12人。"},
        ],
    )
    assert (
        store.add_candidates(
            sid, "first", [candidate()], extraction=(0, "model", "version")
        )
        == 1
    )
    repeated = candidate(quote="研发人数合计12人。")
    assert (
        store.add_candidates(
            sid, "second", [repeated], extraction=(0, "model", "version")
        )
        == 0
    )
    assert len(store.review_rows()) == 1
    assert store.review_rows()[0]["chunk_id"] == "first"
    with store.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM extraction_facts").fetchone()[0] == 2
    repeated["conditions"] = "另一个项目"
    assert (
        store.add_candidates(
            sid, "second", [repeated], extraction=(0, "model", "version")
        )
        == 1
    )
    assert len(store.review_rows()) == 2


def test_extraction_deduplication_keeps_validity_windows_in_context(store):
    sid = source(store)
    for date in ("2025-01-01", "2026-01-01"):
        assert (
            store.add_candidates(
                sid,
                "chunk",
                [candidate(valid_from=date)],
                extraction=(0, "model", "version"),
            )
            == 1
        )
    facts = store.extraction_candidates(sid, "model", "version")
    assert {f["valid_from"] for f in facts} == {"2025-01-01", "2026-01-01"}
    assert len(store.review_rows()) == 2


def test_incomplete_model_dates_are_repaired_without_inventing_day(store, monkeypatch):
    source(store, content="2008.01 任研究员。研发人员为12人。")
    monkeypatch.setenv("LLM_API_KEY", "test")
    calls = []

    def reply(request):
        payload = json.loads(request.content)
        calls.append(payload)
        if len(calls) == 1:
            facts = [candidate(valid_from="2008-01")]
        else:
            correction = json.loads(payload["messages"][-1]["content"])
            assert "Invalid isoformat" in correction["field_errors"]["0"]
            assert "不能补造日" in correction["instruction"]
            facts = [candidate(conditions="任职起始年月：2008.01，未披露具体日期")]
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps({"facts": facts})},
                    }
                ]
            },
        )

    config = {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"}
    assert (
        extract(store, config, httpx.Client(transport=httpx.MockTransport(reply))) == 1
    )
    assert len(calls) == 2
    assert store.review_rows()[0]["valid_from"] == ""
    assert "2008.01" in store.review_rows()[0]["conditions"]
    assert store.verified() == []


def test_unrepaired_fields_never_complete_or_partially_commit_chunk(store, monkeypatch):
    sid = source(store)
    monkeypatch.setenv("LLM_API_KEY", "test")

    def reply(request):
        facts = [candidate(), candidate(valid_until="2027-03")]
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps({"facts": facts})},
                    }
                ]
            },
        )

    with pytest.raises(ValueError, match="invalid fact fields"):
        extract(
            store,
            {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"},
            httpx.Client(transport=httpx.MockTransport(reply)),
        )
    assert store.review_rows() == []
    assert not store.extracted(
        sid, "chunk", 0, "deepseek-flash", "fact-manager-2|document-context-2"
    )


def test_presentation_layout_retains_columns_and_nested_group_coordinates(
    store, monkeypatch
):
    import io
    import zipfile
    from fact_manager.workflow import presentation_layout

    header = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

    def shape(x, y, text):
        return f'<p:sp><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="100" cy="50"/></a:xfrm></p:spPr><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'

    slide = f'<p:sld {header}><p:cSld><p:spTree>{shape(500,100,"Bob")}{shape(100,100,"Alice")}<p:grpSp><p:grpSpPr><a:xfrm><a:off x="100" y="200"/><a:ext cx="200" cy="100"/><a:chOff x="0" y="0"/><a:chExt cx="100" cy="100"/></a:xfrm></p:grpSpPr>{shape(0,0,"Founder")}</p:grpSp></p:spTree></p:cSld></p:sld>'
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        # Presentation order can differ from slide filenames after a user reorders slides.
        archive.writestr(
            "ppt/presentation.xml",
            f'<p:presentation {header} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId r:id="first"/><p:sldId r:id="second"/></p:sldIdLst><p:sldSz cx="1000" cy="500"/></p:presentation>',
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            '<Relationships><Relationship Id="first" Target="slides/slide2.xml"/><Relationship Id="second" Target="slides/slide1.xml"/></Relationships>',
        )
        archive.writestr("ppt/slides/slide2.xml", slide)
        archive.writestr(
            "ppt/slides/slide1.xml",
            f'<p:sld {header}><p:cSld><p:spTree>{shape(0,0,"Wrong page")}</p:spTree></p:cSld></p:sld>',
        )
    sid = store.register(
        "company",
        "slides",
        "团队.pptx",
        buffer.getvalue(),
        [
            {
                "id": "team",
                "content": "Bob\nAlice\nFounder",
                "positions": [[1, 0, 0, 0, 0]],
            },
            {"id": "picture", "content": "", "positions": [[2, 0, 0, 0, 0]]},
        ],
    )
    layout = presentation_layout(store.source(sid), store.chunks(sid)[0])
    assert [(b["text"], b["x"], b["y"], b["width"]) for b in layout["blocks"]] == [
        ("Alice", 10, 20, 10),
        ("Founder", 10, 40, 20),
        ("Bob", 50, 20, 10),
    ]
    assert layout["blocks"][1]["group"]
    monkeypatch.setenv("LLM_API_KEY", "test")
    calls = []

    def reply(request):
        payload = json.loads(request.content)
        calls.append(payload)
        target = json.loads(payload["messages"][1]["content"])
        assert target["target_text"] == "Bob\nAlice\nFounder"
        assert target["target_layout"] == layout
        assert "同栏/同组" in payload["messages"][0]["content"]
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"finish_reason": "stop", "message": {"content": '{"facts": []}'}}
                ]
            },
        )

    config = {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"}
    client = httpx.Client(transport=httpx.MockTransport(reply))
    progress = []
    assert (
        extract(
            store,
            config,
            client,
            progress=lambda done, total: progress.append((done, total)),
        )
        == 0
    )
    assert progress[-1] == (2, 2)
    assert len(calls) == 1  # Empty image-only pages never trigger a model call.
    assert store.extracted(
        sid,
        "picture",
        0,
        "deepseek-flash",
        "fact-manager-2|document-context-2|ppt-layout-2",
    )
    assert extract(store, config, client) == 0
    assert len(calls) == 1
    assert store.chunks(sid)[0]["content"] == "Bob\nAlice\nFounder"


def test_ppt_line_references_preserve_exact_bullets_and_repair_invented_dates(
    store, monkeypatch
):
    text = ".任职时间：2008.01\n.岗位：研究员\n.成果：项目甲"
    store.register(
        "company",
        "ppt",
        "履历.pptx",
        b"not-a-zip",
        [{"id": "page", "content": text, "positions": [[1, 0, 0, 0, 0]]}],
    )
    monkeypatch.setenv("LLM_API_KEY", "test")
    calls = []

    def reply(request):
        payload = json.loads(request.content)
        calls.append(payload)
        supplied = json.loads(payload["messages"][1]["content"])
        assert supplied["target_lines"][0] == {"line": 1, "text": ".任职时间：2008.01"}
        fact = candidate(
            entity="Alice",
            attribute="工作经历",
            value="研究员",
            unit="",
            quote="任职时间：2008.01\n岗位：研究员",
            quote_lines=[1, 2],
            valid_from="2008-01-01" if len(calls) == 1 else "",
            conditions="2008.01起，未披露具体日",
        )
        if len(calls) > 1:
            assert (
                "未披露该完整日期"
                in json.loads(payload["messages"][-1]["content"])["field_errors"]["0"]
            )
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {"content": json.dumps({"facts": [fact]})},
                    }
                ]
            },
        )

    config = {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"}
    assert (
        extract(store, config, httpx.Client(transport=httpx.MockTransport(reply))) == 1
    )
    assert len(calls) == 2
    row = store.review_rows()[0]
    assert row["quote"] == ".任职时间：2008.01\n.岗位：研究员"
    assert row["quote"] in text and row["valid_from"] == ""


def test_invalid_ppt_line_reference_is_rejected_even_with_otherwise_valid_quote(
    store, monkeypatch
):
    store.register(
        "company",
        "ppt",
        "错误.pptx",
        b"not-a-zip",
        [{"id": "page", "content": "研发人员为12人。", "positions": [[1, 0, 0, 0, 0]]}],
    )
    monkeypatch.setenv("LLM_API_KEY", "test")

    def reply(request):
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "stop",
                        "message": {
                            "content": json.dumps(
                                {"facts": [candidate(quote_lines=[1, 999])]}
                            )
                        },
                    }
                ]
            },
        )

    with pytest.raises(ValueError, match="unsupported evidence quote"):
        extract(
            store,
            {"llm_url": "https://api.deepseek.com", "llm_model": "deepseek-flash"},
            httpx.Client(transport=httpx.MockTransport(reply)),
        )
    assert store.review_rows() == []


def test_ppt_control_characters_and_literal_escape_text_roundtrip_without_changing_evidence(
    store, tmp_path
):
    from fact_manager.workflow import EXCEL_TEXT_ENCODING

    text = "场景：千卡集群\x0b\n提升12%。原文中的字面转义：\\u000B和路径C:\\data。"
    sid = source(store, content=text)
    store.add_candidates(
        sid, "chunk", [candidate(value="C:\\data及字面\\u000B", unit="", quote=text)]
    )
    path = tmp_path / "ppt-review.xlsx"
    assert export_review(store, path) == 1
    book = load_workbook(path)
    assert book.properties.description == EXCEL_TEXT_ENCODING
    displayed = book["facts"].cell(2, HEADERS.index("quote") + 1).value
    assert "\x0b" not in displayed
    assert "\\u000B" in displayed and "\\\\u000B" in displayed
    book["facts"].cell(2, HEADERS.index("decision") + 1).value = "approve"
    book.save(path)
    assert import_review(store, path, "审核人") == 1
    with store.connect() as db:
        row = db.execute("SELECT * FROM facts").fetchone()
        assert row["quote"] == text and row["value"] == "C:\\data及字面\\u000B"
    assert store.verified() == []


def test_escaped_control_character_quote_tampering_is_still_rejected(store, tmp_path):
    sid = source(store, content="研发人员为12人。\x0b真实条件。")
    store.add_candidates(
        sid, "chunk", [candidate(quote="研发人员为12人。\x0b真实条件。")]
    )
    path = tmp_path / "tampered-ppt.xlsx"
    export_review(store, path)
    book = load_workbook(path)
    sheet = book["facts"]
    sheet.cell(2, HEADERS.index("quote") + 1).value = "研发人员为12人。 真实条件。"
    sheet.cell(2, HEADERS.index("decision") + 1).value = "approve"
    book.save(path)
    with pytest.raises(ValueError, match="Evidence fields"):
        import_review(store, path, "审核人")
    assert store.review_rows()[0]["status"] == "pending"
